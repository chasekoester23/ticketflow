import io
from decimal import Decimal

from openpyxl import load_workbook

from .sample_data import make_scan_pdf


def _project(client, code="P-100", name="Highway 31"):
    r = client.post("/api/projects", json={"code": code, "name": name, "job_location": "Project A"})
    assert r.status_code == 201
    return r.json()


def _upload(client, project_id=None):
    data = {"layout": "grid_2x2"}
    if project_id is not None:
        data["project_id"] = str(project_id)
    return client.post("/api/batches", data=data, files=[("files", ("scan.pdf", make_scan_pdf(), "application/pdf"))])


def test_project_crud_and_duplicate_code(client):
    p = _project(client)
    assert client.post("/api/projects", json={"code": "P-100", "name": "Other"}).status_code == 409
    assert [x["code"] for x in client.get("/api/projects").json()] == ["P-100"]
    r = client.patch(f"/api/projects/{p['id']}", json={"status": "archived"})
    assert r.json()["status"] == "archived"
    assert client.get("/api/projects").json() == []  # archived hidden by default
    assert len(client.get("/api/projects?include_archived=true").json()) == 1
    assert client.patch(f"/api/projects/{p['id']}", json={"status": "weird"}).status_code == 422
    assert client.patch("/api/projects/999", json={"name": "x"}).status_code == 404


def test_upload_requires_valid_active_project(client):
    assert _upload(client, project_id=999).status_code == 404
    p = _project(client)
    r = _upload(client, p["id"])
    assert r.status_code == 200 and r.json()["project_id"] == p["id"]
    client.patch(f"/api/projects/{p['id']}", json={"status": "archived"})
    assert _upload(client, p["id"]).status_code == 409


def test_records_filters_by_project_supplier_material_status_date(client):
    p1, p2 = _project(client), _project(client, "P-200", "Other job")
    _upload(client, p1["id"])
    _upload(client, p2["id"])
    assert len(client.get("/api/records").json()) == 8
    assert len(client.get(f"/api/records?project_id={p1['id']}").json()) == 4
    assert len(client.get("/api/records?supplier=sample").json()) == 4  # 2 per batch
    assert len(client.get("/api/records?material=sand").json()) == 4
    assert len(client.get("/api/records?date_from=2026-10-16").json()) == 4
    assert len(client.get("/api/records?date_to=2026-10-15").json()) == 4
    assert len(client.get("/api/records?date_from=2026-10-16&date_to=2026-10-15").json()) == 0

    first = client.get(f"/api/records?project_id={p1['id']}").json()[0]
    client.post(f"/api/tickets/{first['id']}/verify")
    assert len(client.get("/api/records?status=verified").json()) == 1
    assert len(client.get("/api/records?status=needs_review").json()) == 7
    client.post(f"/api/tickets/{first['id']}/skip")
    assert len(client.get("/api/records").json()) == 7  # skipped hidden unless asked for
    assert len(client.get("/api/records?status=skipped").json()) == 1


def test_dashboard_totals_with_pricing(client):
    s = client.post("/api/suppliers", json={"name": "Example Aggregates"}).json()
    m = client.post("/api/materials", json={"name": "No. 53 Crushed Stone"}).json()
    client.put("/api/rates", json={"supplier_id": s["id"], "material_id": m["id"], "unit_price": "30"})
    p = _project(client)
    _upload(client, p["id"])
    ts = client.get(f"/api/records?project_id={p['id']}").json()
    client.post(f"/api/tickets/{ts[0]['id']}/verify")

    d = client.get("/api/dashboard").json()
    assert d["tickets"] == 4 and d["verified"] == 1 and d["needs_review"] == 3
    assert Decimal(str(d["total_tons"])) == Decimal("86.65")  # 24.5 + 22.1 + 19.75 + 20.3
    assert Decimal(str(d["total_cost"])) == Decimal("735.00")
    assert d["unpriced_tickets"] == 3
    assert {r["name"] for r in d["by_supplier"]} == {"Example Aggregates", "Sample Materials"}
    assert d["recent_batches"][0]["project"] == "Highway 31"
    assert client.get("/api/dashboard?project_id=999").json()["tickets"] == 0


def test_filtered_export_matches_filter(client):
    p1, p2 = _project(client), _project(client, "P-200", "Other job")
    _upload(client, p1["id"])
    _upload(client, p2["id"])
    r = client.get(f"/api/records/export.xlsx?project_id={p1['id']}&supplier=sample")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Delivery Records"]
    assert ws.max_row == 3  # header + 2 Sample Materials tickets from project 1 only
    assert {ws.cell(i, 5).value for i in (2, 3)} == {"Sample Materials"}
