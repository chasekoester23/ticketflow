"""End-to-end: upload a synthetic 2x2 scan -> split -> (fixture) extract -> review -> verify -> export."""
import io

from openpyxl import load_workbook

from .sample_data import make_scan_pdf


def _upload(client, name="scan.pdf", data=None, layout="grid_2x2"):
    data = data if data is not None else make_scan_pdf()
    return client.post("/api/batches", data={"layout": layout}, files=[("files", (name, data, "application/pdf"))])


def test_full_workflow(client):
    r = _upload(client)
    assert r.status_code == 200
    batch_id = r.json()["id"]

    batch = client.get(f"/api/batches/{batch_id}").json()
    assert batch["status"] == "ready_for_review"
    assert batch["files"][0]["status"] == "ready_for_review"

    tickets = client.get(f"/api/batches/{batch_id}/tickets").json()
    assert len(tickets) == 4
    assert [t["ticket_number"] for t in tickets] == ["10482", "10483", "77120", "77121"]
    assert all(t["status"] == "needs_review" and not t["tonnage_confirmed"] for t in tickets)

    img = client.get(f"/api/tickets/{tickets[0]['id']}/image")
    assert img.status_code == 200 and img.headers["content-type"] == "image/png"

    # Correct a tonnage -> audited, still unconfirmed
    tid = tickets[0]["id"]
    r = client.patch(f"/api/tickets/{tid}", json={"tonnage": "24.75"})
    assert r.status_code == 200 and r.json()["tonnage"] == "24.750"
    audit = client.get(f"/api/tickets/{tid}/audit").json()
    assert audit[0]["field"] == "tonnage" and audit[0]["old"] == "24.500" and audit[0]["new"] == "24.750"

    # Verify two, skip one, leave one for review
    assert client.post(f"/api/tickets/{tickets[0]['id']}/verify").json()["status"] == "verified"
    assert client.post(f"/api/tickets/{tickets[1]['id']}/verify").json()["tonnage_confirmed"] is True
    assert client.post(f"/api/tickets/{tickets[2]['id']}/skip").json()["status"] == "skipped"

    # Editing a verified ticket knocks it back to review (no silent changes)
    r = client.patch(f"/api/tickets/{tickets[1]['id']}", json={"tonnage": "22.20"}).json()
    assert r["status"] == "needs_review" and r["tonnage_confirmed"] is False

    xlsx = client.get(f"/api/batches/{batch_id}/export.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Delivery Records", "Summary"]
    ws = wb["Delivery Records"]
    assert [c.value for c in ws[1]][:6] == ["Ticket Number", "Delivery Date", "Material Type", "Tonnage", "Supplier", "Job Location"]
    assert ws.max_row == 4  # header + 3 non-skipped tickets
    assert ws["A2"].value == "10482" and ws["D2"].value == 24.75 and ws["I2"].value == "Verified"
    assert ws["H2"].value is None  # no rate configured yet -> cost stays unresolved, never invented
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref


def test_verify_requires_tonnage(client):
    batch_id = _upload(client).json()["id"]
    t = client.get(f"/api/batches/{batch_id}/tickets").json()[0]
    client.patch(f"/api/tickets/{t['id']}", json={"tonnage": None})
    assert client.post(f"/api/tickets/{t['id']}/verify").status_code == 422


def test_rejects_bad_uploads(client):
    assert _upload(client, name="notes.txt", data=b"hello").status_code == 400
    assert _upload(client, name="empty.pdf", data=b"").status_code == 400
    assert _upload(client, name="fake.pdf", data=b"MZ not a pdf").status_code == 400
    assert _upload(client, layout="bogus").status_code == 422


def test_corrupt_pdf_marks_file_failed_not_crash(client):
    r = _upload(client, name="bad.pdf", data=b"%PDF-1.4\ngarbage")
    batch = client.get(f"/api/batches/{r.json()['id']}").json()
    assert batch["status"] == "failed"
    assert batch["files"][0]["status"] == "failed" and batch["files"][0]["error"]
