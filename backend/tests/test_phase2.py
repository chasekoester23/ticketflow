import io
from decimal import Decimal

from openpyxl import load_workbook

from .sample_data import TICKETS, make_scan_pdf
from .test_api_e2e import _upload


def _setup_rates(client):
    s = client.post("/api/suppliers", json={"name": "Example Aggregates"}).json()
    m = client.post("/api/materials", json={"name": "No. 53 Crushed Stone"}).json()
    r = client.put("/api/rates", json={"supplier_id": s["id"], "material_id": m["id"], "unit_price": "30"})
    assert r.status_code == 200
    return s, m


def test_rate_crud_and_validation(client):
    s, m = _setup_rates(client)
    assert client.post("/api/suppliers", json={"name": "Example Aggregates"}).status_code == 409
    assert client.put("/api/rates", json={"supplier_id": s["id"], "material_id": m["id"], "unit_price": "0"}).status_code == 422
    assert client.put("/api/rates", json={"supplier_id": 999, "material_id": m["id"], "unit_price": "5"}).status_code == 404
    # upsert replaces, not duplicates
    client.put("/api/rates", json={"supplier_id": s["id"], "material_id": m["id"], "unit_price": "31.5"})
    rates = client.get("/api/rates").json()
    assert len(rates) == 1 and Decimal(rates[0]["unit_price"]) == Decimal("31.5")


def test_cost_requires_confirmation_and_matching_rate(client):
    _setup_rates(client)
    bid = _upload(client).json()["id"]
    ts = client.get(f"/api/batches/{bid}/tickets").json()
    first = ts[0]
    assert Decimal(first["unit_price"]) == Decimal("30") and first["price_source"] == "rate"
    assert first["cost"] is None  # tonnage not yet confirmed
    v = client.post(f"/api/tickets/{first['id']}/verify").json()
    assert Decimal(v["cost"]) == Decimal("735.00")  # 24.5 x 30
    # a ticket with no rate stays unresolved and explains why
    other = ts[2]
    assert other["unit_price"] is None and other["cost"] is None
    assert any("Sample Materials" in n for n in other["pricing_notes"])


def test_override_and_export_costs(client):
    _setup_rates(client)
    bid = _upload(client).json()["id"]
    ts = client.get(f"/api/batches/{bid}/tickets").json()
    t = ts[2]  # no configured rate
    r = client.patch(f"/api/tickets/{t['id']}", json={"unit_price_override": "28"}).json()
    assert r["price_source"] == "override" and Decimal(r["unit_price"]) == Decimal("28.0000")
    client.post(f"/api/tickets/{t['id']}/verify")
    client.post(f"/api/tickets/{ts[0]['id']}/verify")
    wb = load_workbook(io.BytesIO(client.get(f"/api/batches/{bid}/export.xlsx").content))
    ws = wb["Delivery Records"]
    assert ws["G2"].value == 30 and ws["H2"].value == 735.0
    assert ws["H4"].value == 553.0  # 19.75 x 28
    assert ws["H3"].value is None  # unpriced / unconfirmed stays blank
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=3, max_row=7)}
    assert summary["Total calculated material cost"] == 735.0 + 553.0
    assert summary["Records without a calculated cost"] == 2


def test_duplicate_same_supplier_and_number_is_flagged_not_deleted(client):
    _upload(client)
    b2 = _upload(client).json()["id"]  # identical scan uploaded again
    ts = client.get(f"/api/batches/{b2}/tickets").json()
    assert len(ts) == 4
    assert all(any("Identical scanned image" in d for d in t["duplicates"]) for t in ts)


def test_same_number_different_supplier_not_flagged(client, monkeypatch):
    from app.routers import api as api_module
    from .conftest import FixtureExtractor

    _upload(client, data=make_scan_pdf(TICKETS[:1]))  # supplier "Example Aggregates", ticket 10482
    clash = dict(TICKETS[1], ticket_number="10482", supplier="Other Quarry")
    monkeypatch.setattr(api_module, "get_extractor", lambda: FixtureExtractor([clash]))
    b2 = _upload(client, data=make_scan_pdf(TICKETS[1:2])).json()["id"]
    t = client.get(f"/api/batches/{b2}/tickets").json()[0]
    assert t["ticket_number"] == "10482" and t["supplier"] == "Other Quarry"
    assert not any("ticket number" in d for d in t["duplicates"])


def test_same_supplier_and_number_flagged(client, monkeypatch):
    from app.routers import api as api_module
    from .conftest import FixtureExtractor

    _upload(client, data=make_scan_pdf(TICKETS[:1]))
    clash = dict(TICKETS[1], ticket_number="10482")  # same supplier, same number, different scan
    monkeypatch.setattr(api_module, "get_extractor", lambda: FixtureExtractor([clash]))
    b2 = _upload(client, data=make_scan_pdf(TICKETS[1:2])).json()["id"]
    t = client.get(f"/api/batches/{b2}/tickets").json()[0]
    assert any("Same supplier and ticket number 10482" in d for d in t["duplicates"])


def test_status_endpoint_hides_key(client):
    body = client.get("/api/status").json()
    assert set(body) == {"extraction_configured", "model"}
